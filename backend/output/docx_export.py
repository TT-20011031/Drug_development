import os
import re
import uuid
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "generated_docs")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def _add_markdown_table(doc, table_lines):
    rows = []
    for line in table_lines:
        line = line.strip()
        if line.startswith("|") and not re.match(r"^\|[\s\-:|]+\|$", line):
            cells = [c.strip() for c in line.split("|")[1:-1]]
            rows.append(cells)

    if len(rows) < 1:
        return

    header = rows[0]
    data_rows = rows[1:]

    table = doc.add_table(rows=1 + len(data_rows), cols=len(header))
    table.style = "Table Grid"

    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)

    for row_idx, row_data in enumerate(data_rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx < len(header):
                cell = table.rows[row_idx + 1].cells[col_idx]
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)


def markdown_to_docx(markdown_text: str, filename: str = None) -> str:
    if not filename:
        filename = f"product_dev_{uuid.uuid4().hex[:8]}.docx"
    if not filename.endswith(".docx"):
        filename += ".docx"

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "微软雅黑"
    font.size = Pt(11)

    lines = markdown_text.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            heading = doc.add_heading(stripped[2:], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("").paragraph_format.space_after = Pt(12)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _add_markdown_table(doc, table_lines)
            doc.add_paragraph("")
            continue

        if stripped.startswith("- **") or stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            text = stripped.lstrip("- ").strip()
            bold_match = re.match(r"\*\*(.+?)\*\*[：:]\s*(.*)", text)
            if bold_match:
                run = p.add_run(bold_match.group(1) + "：")
                run.bold = True
                run.font.size = Pt(11)
                run = p.add_run(bold_match.group(2))
                run.font.size = Pt(11)
            else:
                text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
                run = p.add_run(text)
                run.font.size = Pt(11)
            i += 1
            continue

        text = re.sub(r"\*\*(.+?)\*\*", r"\1", stripped)
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(11)
        i += 1

    filepath = os.path.join(OUTPUT_DIR, filename)
    doc.save(filepath)
    return filepath
